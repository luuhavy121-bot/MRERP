"""Bounded reader for the HR 'Chi Tiet' export. Never calculates attendance."""
import re
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation
from io import BytesIO
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from rest_framework.exceptions import ValidationError

HEADER = re.compile(r"Mã nhân viên:\s*(.*?)\s+Tên nhân viên:\s*(.*?)\s+Bộ phận:", re.I)
METRICS = ("late", "early", "workdays", "hours", "overtime1", "overtime2", "overtime3")
MAX_BYTES = 10 * 1024 * 1024


def parse_workbook(content):
    if len(content) > MAX_BYTES:
        raise ValidationError("File vượt quá 10 MB.")
    try:
        with ZipFile(BytesIO(content)) as archive:
            if len(archive.infolist()) > 1000 or sum(i.file_size for i in archive.infolist()) > 50 * 1024 * 1024:
                raise ValidationError("File giải nén quá lớn.")
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
    except (BadZipFile, ValueError, KeyError, OSError) as exc:
        raise ValidationError("Không đọc được file Excel .xlsx.") from exc
    try:
        if "Chi Tiet" not in workbook.sheetnames:
            raise ValidationError("Không tìm thấy sheet Chi Tiet của mẫu HR.")
        sheet = workbook["Chi Tiet"]
        if (sheet.max_row or 0) > 20000 or (sheet.max_column or 0) > 64:
            raise ValidationError("Bảng vượt giới hạn 20.000 dòng hoặc sai định dạng.")
        employees, errors, months, seen = [], [], set(), set()
        current = None
        ready = False
        for rowno, cells in enumerate(sheet.iter_rows(max_col=16), 1):
            if rowno > 20000:
                raise ValidationError("Bảng vượt giới hạn 20.000 dòng.")
            values = [c.value for c in cells]
            first = values[0]
            if first == "BẢNG CHI TIẾT CHẤM CÔNG":
                current, ready = None, False
            elif isinstance(first, str) and "Mã nhân viên:" in first:
                match = HEADER.search(first)
                if not match or not match[1].strip():
                    raise ValidationError(f"Dòng {rowno}: thiếu mã hoặc tên nhân viên.")
                code, name = match[1].strip(), match[2].strip()
                if len(code) > 64 or any(e["code"] == code for e in employees):
                    raise ValidationError(f"Dòng {rowno}: mã quá dài hoặc bị lặp trong file.")
                current = {"code": code, "name": name[:200], "rows": []}
                employees.append(current)
                ready = False
            elif first == "Ngày":
                ready = values[8:15] == ["Trễ", "Sớm", "Công", "T.Giờ", "T.Ca1", "T.Ca2", "T.Ca3"]
                if not ready:
                    raise ValidationError(f"Dòng {rowno}: cột chấm công không đúng mẫu.")
            elif isinstance(first, (datetime, date)):
                if not current or not ready:
                    raise ValidationError(f"Dòng {rowno}: không xác định được nhân sự hoặc tiêu đề.")
                day = first.date() if isinstance(first, datetime) else first
                months.add(day.strftime("%Y-%m"))
                key = (current["code"], day.isoformat())
                if key in seen:
                    errors.append(f"Dòng {rowno}: ngày chấm công bị lặp.")
                    continue
                seen.add(key)
                parsed = {"date": day.isoformat(), "source_row": rowno, "punches": []}
                for value in values[2:8]:
                    if isinstance(value, (datetime, time)):
                        value = value.strftime("%H:%M:%S")
                    parsed["punches"].append(str(value).strip()[:80] if value not in (None, "") else None)
                for field, value in zip(METRICS, values[8:15]):
                    if value in (None, ""):
                        parsed[field] = None
                        continue
                    try:
                        number = Decimal(str(value))
                        if not number.is_finite() or number < 0 or number > 100000 or number.as_tuple().exponent < -8:
                            raise InvalidOperation
                        parsed[field] = str(number)
                    except InvalidOperation:
                        errors.append(f"Dòng {rowno}: giá trị cột {field} không hợp lệ.")
                        parsed[field] = None
                if any(c.data_type == "f" for c in cells):
                    errors.append(f"Dòng {rowno}: chỉ nhận kết quả xuất, không nhận công thức Excel.")
                current["rows"].append(parsed)
            elif current and ready and first not in (None, ""):
                errors.append(f"Dòng {rowno}: ngày không đúng định dạng ngày Excel.")
        if len(employees) > 500:
            raise ValidationError("Mỗi file tối đa 500 nhân sự.")
        if not seen:
            raise ValidationError("File chưa có dòng chấm công. Hãy xuất bảng chi tiết có ngày từ phần mềm HR.")
        if len(months) != 1:
            raise ValidationError("Mỗi lần nhập chỉ nhận một tháng chấm công.")
        if any(not e["rows"] for e in employees):
            errors.append("Có nhân sự không có dòng chấm công.")
        return {"month": next(iter(months)), "employees": employees, "errors": errors[:100], "row_count": len(seen)}
    finally:
        workbook.close()
