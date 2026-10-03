from datetime import datetime
from io import BytesIO

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Permission
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from openpyxl import Workbook
from rest_framework.test import APITestCase

from people_domain.models import AuditEvent, Employee
from .attendance_parser import parse_workbook
from .models import AttendanceEmployeeMapping, AttendanceImport, AttendanceRecord


def fixture_bytes(workdays="0.5", code="00001", second=False, formula=False):
    """Synthetic workbook, never derived from real employee attendance."""
    book = Workbook()
    sheet = book.active
    sheet.title = "Chi Tiet"
    for source_code in ([code, "00002"] if second else [code]):
        sheet.append(["BẢNG CHI TIẾT CHẤM CÔNG"])
        sheet.append([f"Mã nhân viên: {source_code}   Tên nhân viên: Nhân sự mẫu   Bộ phận: ---"])
        sheet.append(["Ngày", "Thứ", "1", None, "2", None, "3", None, "Trễ", "Sớm", "Công", "T.Giờ", "T.Ca1", "T.Ca2", "T.Ca3", "Nơi làm việc"])
        sheet.append([datetime(2026, 9, 7), "Hai", "08:00", "12:00", None, None, None, None, 0, 1, "=1/2" if formula else workdays, "3.48", 0, 0, 0])
        sheet.append([datetime(2026, 9, 8), "Ba", None, None, None, None, None, None, 0, 0, 0, None, 0, 0, 0])
    stream = BytesIO()
    book.save(stream)
    return stream.getvalue()


class AttendanceImportTests(APITestCase):
    @classmethod
    def setUpTestData(cls):
        call_command("seed_demo", password="Test-Only-1234!", verbosity=0)
        for attr, name in [("hr", "hr.demo"), ("staff", "staff.demo"), ("other", "other.demo"), ("leader", "leader.demo"), ("ceo", "ceo.demo")]:
            setattr(cls, attr, get_user_model().objects.get(username=name))

    def preview(self, user=None, content=None):
        self.client.force_authenticate(user or self.hr)
        return self.client.post("/api/v1/leave/attendance-imports/preview/", {"file": SimpleUploadedFile("synthetic.xlsx", content if content is not None else fixture_bytes())}, format="multipart")

    def commit(self, preview, mappings=None, replace=False):
        return self.client.post(f"/api/v1/leave/attendance-imports/{preview.data['uuid']}/commit/",
                                {"mappings": mappings or {"00001": str(self.staff.employee_profile.pk)}, "replace_existing": replace}, format="json")

    def listing(self, user):
        self.client.force_authenticate(user)
        return self.client.get("/api/v1/leave/actual-attendance/?month=2026-09")

    def test_parser_preserves_codes_zero_blank_and_source_results(self):
        data = parse_workbook(fixture_bytes())
        self.assertEqual(data["employees"][0]["code"], "00001")
        rows = data["employees"][0]["rows"]
        self.assertEqual(rows[0]["workdays"], "0.5")
        self.assertEqual(rows[0]["hours"], "3.48")
        self.assertEqual(rows[1]["workdays"], "0")
        self.assertIsNone(rows[1]["hours"])

    def test_preview_does_not_write_records_and_commit_is_idempotent(self):
        preview = self.preview()
        self.assertEqual(preview.status_code, 201, preview.data)
        self.assertEqual(AttendanceRecord.objects.count(), 0)
        response = self.commit(preview)
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(self.commit(preview).status_code, 200)
        self.assertEqual(AttendanceRecord.objects.count(), 2)
        self.assertEqual(AuditEvent.objects.filter(action="attendance.import.committed").count(), 1)
        self.assertEqual(AttendanceEmployeeMapping.objects.get().source_code, "00001")
        again = self.preview()
        self.assertEqual(again.data["mappings"]["00001"], str(self.staff.employee_profile.pk))

    def test_replacement_explicit_and_history_preserved(self):
        first = self.preview()
        self.commit(first)
        second = self.preview(content=fixture_bytes(workdays="1"))
        self.assertEqual(self.commit(second).status_code, 409)
        self.assertEqual(self.commit(second, replace=True).status_code, 200)
        self.assertEqual(AttendanceRecord.objects.count(), 2)
        self.assertEqual(AttendanceRecord.objects.first().data["workdays"], "1")
        self.assertEqual(AttendanceImport.objects.get(pk=first.data["uuid"]).source["employees"][0]["rows"][0]["workdays"], "0.5")

    def test_two_previews_cannot_overwrite_without_refresh(self):
        first, second = self.preview(), self.preview()
        self.assertEqual(self.commit(first).status_code, 200)
        self.assertEqual(self.commit(second, replace=True).status_code, 409)

    def test_staff_and_leader_scope_and_field_redaction(self):
        preview = self.preview(content=fixture_bytes(second=True))
        self.assertEqual(self.commit(preview, {"00001": str(self.staff.employee_profile.pk), "00002": str(self.other.employee_profile.pk)}).status_code, 200)
        for user in (self.staff, self.leader):
            response = self.listing(user)
            self.assertEqual(response.status_code, 200)
            self.assertEqual([e["code"] for e in response.data["employees"]], ["STF01"])
            self.assertNotIn("source", response.data)
            self.assertNotIn("Nhân sự mẫu", str(response.data))
            self.assertNotIn("imported_by", str(response.data))
            self.assertNotIn("address", str(response.data))
        self.assertEqual(len(self.listing(self.hr).data["employees"]), 2)
        self.assertEqual(len(self.listing(self.ceo).data["employees"]), 2)

    def test_import_and_history_denied_without_capability(self):
        for user in (self.staff, self.leader, self.ceo):
            self.assertEqual(self.preview(user=user).status_code, 403)
            self.assertEqual(self.client.get("/api/v1/leave/attendance-imports/").status_code, 403)

    def test_missing_read_capability_denied(self):
        self.staff.groups.clear()
        self.assertEqual(self.listing(self.staff).status_code, 403)

    def test_invalid_account_or_employment_denied(self):
        for active, employment in [(False, "official"), (True, "terminated"), (True, "paused")]:
            self.hr.is_active = active
            self.hr.save()
            employee = self.hr.employee_profile
            employee.employment_status = employment
            employee.save()
            self.assertEqual(self.preview().status_code, 403)
            self.assertEqual(self.listing(self.hr).status_code, 403)
            self.assertEqual(self.client.get("/api/v1/leave/attendance-imports/").status_code, 403)

    def test_cannot_commit_another_hr_preview(self):
        preview = self.preview()
        self.other.user_permissions.add(Permission.objects.get(codename="import_attendance"))
        self.client.force_authenticate(self.other)
        self.assertEqual(self.commit(preview).status_code, 404)

    def test_unmapped_duplicate_or_unknown_employee_rolls_back(self):
        preview = self.preview(content=fixture_bytes(second=True))
        self.assertEqual(self.commit(preview).status_code, 400)
        self.assertEqual(self.commit(preview, {"00001": str(self.staff.employee_profile.pk), "00002": str(self.staff.employee_profile.pk)}).status_code, 400)
        self.assertEqual(AttendanceRecord.objects.count(), 0)
        self.assertEqual(AttendanceEmployeeMapping.objects.count(), 0)

    def test_saved_mapping_cannot_be_reassigned(self):
        self.commit(self.preview())
        preview = self.preview()
        self.assertEqual(self.commit(preview, {"00001": str(self.other.employee_profile.pk)}, True).status_code, 409)

    def test_formula_negative_and_non_numeric_errors_block_commit(self):
        for content in [fixture_bytes(formula=True), fixture_bytes(workdays="-1"), fixture_bytes(workdays="wrong")]:
            preview = self.preview(content=content)
            self.assertEqual(preview.status_code, 201)
            self.assertTrue(preview.data["source"]["errors"])
            self.assertEqual(self.commit(preview).status_code, 400)

    def test_corrupt_wrong_type_empty_and_oversize_rejected(self):
        self.assertEqual(self.preview(content=b"garbage").status_code, 400)
        self.assertEqual(self.preview(content=b"x" * (10 * 1024 * 1024 + 1)).status_code, 400)
        book = Workbook(); book.active.title = "Chi Tiet"
        stream = BytesIO(); book.save(stream)
        self.assertEqual(self.preview(content=stream.getvalue()).status_code, 400)

    def test_no_automatic_employee_creation_or_projected_adjustment(self):
        count = Employee.objects.count()
        self.commit(self.preview())
        self.assertEqual(Employee.objects.count(), count)
        actual = self.listing(self.staff).data["employees"][0]
        self.assertEqual(actual["totals"]["workdays"], "0.5")
        self.assertTrue(actual["incomplete"])

    def test_duplicate_dates_codes_wrong_header_and_mixed_months(self):
        from openpyxl import load_workbook
        for kind in ("duplicate_date", "duplicate_code", "header", "month"):
            book = load_workbook(BytesIO(fixture_bytes(second=True)))
            sheet = book.active
            if kind == "duplicate_date":
                sheet["A5"] = sheet["A4"].value
            elif kind == "duplicate_code":
                sheet["A7"] = sheet["A2"].value
            elif kind == "header":
                sheet["K3"] = "Sai cột"
            else:
                sheet["A5"] = datetime(2026, 10, 1)
            stream = BytesIO(); book.save(stream)
            result = self.preview(content=stream.getvalue())
            if result.status_code == 201:
                self.assertTrue(result.data["source"]["errors"])
                self.assertEqual(self.commit(result).status_code, 400)
            else:
                self.assertEqual(result.status_code, 400)


from concurrent.futures import ThreadPoolExecutor
from threading import Barrier
from unittest import skipUnless
from django.db import connection, connections
from django.test import TransactionTestCase
from rest_framework.test import APIClient


@skipUnless(connection.vendor == "postgresql", "Row-lock concurrency requires PostgreSQL")
class AttendanceConcurrencyTests(TransactionTestCase):
    def test_two_imports_only_one_can_write_same_employee_days(self):
        call_command("seed_demo", password="Test-Only-1234!", verbosity=0)
        user = get_user_model().objects.get(username="hr.demo")
        employee = Employee.objects.get(employee_code="STF01")
        client = APIClient(); client.force_authenticate(user)
        batches = []
        for _ in range(2):
            response = client.post("/api/v1/leave/attendance-imports/preview/", {"file": SimpleUploadedFile("synthetic.xlsx", fixture_bytes())}, format="multipart")
            self.assertEqual(response.status_code, 201)
            batches.append(response.data["uuid"])
        barrier = Barrier(2)
        def commit(uuid):
            try:
                api = APIClient(); api.force_authenticate(get_user_model().objects.get(pk=user.pk))
                barrier.wait(timeout=10)
                return api.post(f"/api/v1/leave/attendance-imports/{uuid}/commit/", {"mappings": {"00001": str(employee.pk)}, "replace_existing": True}, format="json").status_code
            finally:
                connections.close_all()
        with ThreadPoolExecutor(max_workers=2) as pool:
            codes = list(pool.map(commit, batches))
        self.assertEqual(sorted(codes), [200, 409])
        self.assertEqual(AttendanceRecord.objects.count(), 2)
        self.assertEqual(AttendanceImport.objects.filter(committed_at__isnull=False).count(), 1)
